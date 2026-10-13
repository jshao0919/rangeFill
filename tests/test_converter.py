"""Numerical safeguards and regressions against the user's independently entered scores."""
import importlib.util
import json
import os
from pathlib import Path

import cv2
import fitz
import numpy as np
from openpyxl import load_workbook
import pytest

from pdf_to_xlsx import (DISTANCES, Lane, ScanError, Sheet, check_lane_order,
                         collect_inputs, extract_page, find_grid, read_bit,
                         main, validate_lane, write_workbook)


def glyph(text):
    image = np.full((40, 70), 255, np.uint8)
    if text == "0":
        cv2.ellipse(image, (35, 21), (6, 11), 0, 0, 360, 0, 2, cv2.LINE_AA)
    elif text == "1":
        cv2.polylines(image, [np.array([[31, 13], [36, 10], [36, 32]], np.int32)], False, 0, 2, cv2.LINE_AA)
    else:
        cv2.putText(image, text, (25, 30), cv2.FONT_HERSHEY_SIMPLEX, .8, 0, 2, cv2.LINE_AA)
    return image


@pytest.mark.parametrize("digit", [0, 1])
def test_binary_recognition(digit):
    assert read_bit(glyph(str(digit))) == digit


def test_missing_and_nonbinary_glyphs_stay_blank():
    assert read_bit(np.full((40, 70), 255, np.uint8)) is None
    assert read_bit(glyph("2")) is None


def test_total_validation_never_changes_hits():
    hits = [1] * 24 + [0] * 16
    lane = Lane("L1", hits.copy(), 25, glyph("1"), glyph("25"))
    validate_lane(lane)
    assert lane.status == "REVIEW"
    assert lane.calculated_total == 24
    assert lane.difference == -1
    assert lane.hits == hits


def test_missing_hit_is_not_counted_as_zero():
    lane = Lane("L1", [None] + [0] * 39, 0, glyph("1"), glyph("0"))
    validate_lane(lane)
    assert lane.calculated_total is None
    assert lane.difference is None
    assert lane.status == "REVIEW"


def test_lane_gaps_are_preserved_and_conflicts_are_rejected():
    assert check_lane_order([1, 2, None, 4, 7, 10]) == [1, 2, None, 4, 7, 10]
    assert check_lane_order([1, 2, 17, 4, 5]) == [1, 2, None, 4, 5]
    assert check_lane_order([1, 2, 2, 4]) == [1, None, None, 4]


def test_incomplete_grid_is_rejected():
    image = np.full((900, 1200), 255, np.uint8)
    for y in range(100, 700, 40):
        cv2.line(image, (60, y), (1150, y), 0, 2)
    with pytest.raises(ScanError):
        find_grid(image)


def test_input_discovery_keeps_uppercase_pdfs_and_deduplicates(tmp_path):
    first = tmp_path / "first.PDF"
    first.touch()
    (tmp_path / "notes.txt").touch()
    assert collect_inputs([tmp_path, first]) == [first.resolve()]


def test_workbook_retains_blanks_totals_and_review_images(tmp_path):
    lane = Lane("L1", [1] * 24 + [0] * 16, 24, glyph("1"), glyph("24"))
    unknown = Lane(None, [None] + [0] * 39, None, glyph("2"), glyph("0"))
    for item in (lane, unknown):
        validate_lane(item)
    output = tmp_path / "scores.xlsx"
    write_workbook([Sheet("scan.pdf", 1, [lane, unknown], list(DISTANCES), "07 Oct 2026", 2)], output, [])
    book = load_workbook(output)
    data = book.worksheets[0]
    assert data["G2"].value == "L1"
    assert [data.cell(2, c).value for c in range(8, 48)] == lane.hits
    assert [data.cell(2, c).value for c in range(48, 51)] == [24, 24, 0]
    assert data["H3"].value is None
    assert data["AV3"].value is None
    assert data["AY3"].value == "REVIEW"
    assert len(book["Scan Review"]._images) == 4
    assert data["AD1"].value == "23 · 200 m"
    original = output.read_bytes()
    with pytest.raises(FileExistsError):
        write_workbook([], output, [])
    assert output.read_bytes() == original


@pytest.fixture(scope="session")
def sample_scan():
    path = os.environ.get("RANGEFILL_SAMPLE_PDF")
    if not path:
        pytest.skip("Set RANGEFILL_SAMPLE_PDF to the supplied range24.pdf for scan regressions")
    with fitz.open(path) as document:
        pixmap = document[0].get_pixmap(matrix=fitz.Matrix(3, 3), colorspace=fitz.csGRAY, alpha=False)
        return np.frombuffer(pixmap.samples, np.uint8).reshape(pixmap.height, pixmap.width).copy()


def transformed(image, name):
    h, w = image.shape
    if name.startswith("rotate"):
        angle = int(name.removeprefix("rotate"))
        matrix = cv2.getRotationMatrix2D((w / 2, h / 2), angle, 1)
        new_w = int(np.ceil(abs(matrix[0, 0]) * w + abs(matrix[0, 1]) * h))
        new_h = int(np.ceil(abs(matrix[1, 0]) * w + abs(matrix[1, 1]) * h))
        matrix[:, 2] += ((new_w - w) / 2, (new_h - h) / 2)
        return cv2.warpAffine(image, matrix, (new_w, new_h), borderValue=255)
    if name == "perspective":
        src = np.float32([[0, 0], [w - 1, 0], [w - 1, h - 1], [0, h - 1]])
        dst = np.float32([[170, 90], [w - 100, 25], [w - 30, h - 70], [45, h - 20]])
        return cv2.warpPerspective(image, cv2.getPerspectiveTransform(src, dst), (w, h), borderValue=255)
    if name == "quarter_turn":
        return cv2.rotate(image, cv2.ROTATE_90_CLOCKWISE)
    if name == "upside_down":
        return cv2.rotate(image, cv2.ROTATE_180)
    return image


@pytest.mark.parametrize("variant", ["original", "rotate6", "rotate-7", "perspective", "quarter_turn", "upside_down"])
def test_scan_values_and_critical_lane_survive_angle_changes(sample_scan, variant, tmp_path):
    reference = json.loads((Path(__file__).parent / "fixtures/range24_expected.json").read_text())
    sheet = extract_page(transformed(sample_scan, variant), variant + ".pdf", 1, tmp_path)
    assert len(sheet.lanes) == 16
    for index, lane in enumerate(sheet.lanes):
        label = f"L{index + 1}"
        expected = list(map(int, reference["hits"][label]))
        assert lane.label is None or lane.label == label
        # Unreadable cells must be flagged; any accepted digit must be correct.
        for actual, required in zip(lane.hits, expected):
            assert actual is None or actual == required
        if lane.label is None or None in lane.hits or lane.printed_total != reference["totals"][index]:
            assert lane.status == "REVIEW"
        if variant == "original":
            assert lane.status == "MATCH"
            assert lane.hits == expected
            assert lane.printed_total == reference["totals"][index]
    critical = sheet.lanes[0]
    assert critical.label == "L1"
    assert critical.hits == list(map(int, reference["hits"]["L1"]))
    assert critical.calculated_total == critical.printed_total == 24
    assert critical.status == "MATCH"


def test_existing_form_reader_accepts_output(sample_scan, tmp_path):
    reader_path = os.environ.get("RANGEFILL_FORM_READER")
    if not reader_path:
        pytest.skip("Set RANGEFILL_FORM_READER to the supplied xlsx_toForm.py")
    spec = importlib.util.spec_from_file_location("existing_form_reader", reader_path)
    reader = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(reader)
    sheet = extract_page(sample_scan, "range24.pdf", 1)
    output = tmp_path / "compatible.xlsx"
    write_workbook([sheet], output, [])
    fields = {f"stage1_phase{phase}_hit_{i:02d}" for phase in range(1, 5) for i in range(1, 11)}
    fields |= {"lane_firing_order", "table_vi", "date_yyyymmdd"}
    records = list(reader.make_transposed_records(reader.xlsx_rows(output, None), fields))
    assert len(records) == 16
    label, first = records[0]
    assert label == "L1"
    bits = [first[f"stage1_phase{phase}_hit_{i:02d}"] for phase in range(1, 5) for i in range(1, 11)]
    assert sum(int(value) for value in bits) == 24
    assert "Lane 1" in first["lane_firing_order"]


def test_multipage_pdf_keeps_scores_and_reports_an_unreadable_page(sample_scan, tmp_path):
    pdf = tmp_path / "three_pages.pdf"
    with fitz.open() as document:
        for image in (sample_scan, transformed(sample_scan, "rotate6")):
            h, w = image.shape
            page = document.new_page(width=w / 3, height=h / 3)
            page.insert_image(page.rect, stream=cv2.imencode(".png", image)[1].tobytes())
        document.new_page(width=1000, height=760)  # No table: must be reported, not ignored.
        document.save(pdf)
    output = tmp_path / "batch.xlsx"
    assert main([str(pdf), "-o", str(output)]) == 2
    workbook = load_workbook(output)
    data = workbook.worksheets[0]
    assert data.max_row == 33  # 16 lanes per valid page, plus header.
    critical_rows = [row for row in data.iter_rows(min_row=2) if row[6].value == "L1"]
    assert len(critical_rows) == 2
    reference = json.loads((Path(__file__).parent / "fixtures/range24_expected.json").read_text())
    for row in critical_rows:
        assert [cell.value for cell in row[7:47]] == list(map(int, reference["hits"]["L1"]))
        assert [cell.value for cell in row[47:50]] == [24, 24, 0]
    review = workbook["Scan Review"]
    assert review.cell(review.max_row, 2).value == 3
    assert review.cell(review.max_row, 8).value == "ERROR"
