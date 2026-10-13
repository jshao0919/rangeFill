#!/usr/bin/env python3
"""Convert scanned Rifle/Carbine Table VI roll-up sheets to Excel, locally.

python pdf_to_xlsx.py range24.pdf -o range24.xlsx
python pdf_to_xlsx.py scans/ -o scores.xlsx --debug-dir outputs/debug

Requires the Tesseract executable, plus packages in requirements.txt.
Exit codes: 0 = numeric checks passed, 2 = workbook written with review items,
1 = input/runtime failure. Never infer hits from a printed total.
"""
from __future__ import annotations

import argparse
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from io import BytesIO
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys

import cv2
import fitz
import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from scipy.ndimage import gaussian_filter1d
from scipy.signal import find_peaks

TITLE = "RIFLE / CARBINE TABLE VI — RANGE QUALIFICATION"
ROW_HEIGHT = 50
DISTANCES = (50, 100, 150, 50, 150, 200, 150, 200, 250, 300,
             100, 150, 300, 200, 300, 250, 300, 150, 250, 300,
             50, 100, 200, 50, 200, 150, 250, 100, 150, 200,
             50, 100, 200, 250, 50, 100, 150, 100, 200, 250)
cv2.setNumThreads(1)


class ScanError(ValueError):
    """The sheet cannot be located reliably; do not export invented rows."""


class ConverterParser(argparse.ArgumentParser):
    def error(self, message):
        self.print_usage(sys.stderr)
        self.exit(1, f"{self.prog}: error: {message}\n")


@dataclass
class Lane:
    label: str | None
    hits: list[int | None]
    printed_total: int | None
    header_crop: np.ndarray = field(repr=False)
    total_crop: np.ndarray = field(repr=False)
    problems: list[str] = field(default_factory=list)

    @property
    def calculated_total(self) -> int | None:
        return sum(self.hits) if all(v is not None for v in self.hits) else None

    @property
    def difference(self) -> int | None:
        if self.calculated_total is None or self.printed_total is None:
            return None
        return self.calculated_total - self.printed_total

    @property
    def status(self) -> str:
        return "MATCH" if not self.problems else "REVIEW"


@dataclass
class Sheet:
    source: str
    page: int
    lanes: list[Lane]
    distances: list[int | None]
    session: str = ""
    firing_order: int | None = None
    problems: list[str] = field(default_factory=list)


def binary(image: np.ndarray) -> np.ndarray:
    return cv2.threshold(image, 0, 255, cv2.THRESH_BINARY_INV | cv2.THRESH_OTSU)[1]


def hough(image: np.ndarray) -> np.ndarray:
    """Long rules, rather than handwriting, determine page orientation."""
    scale = min(1., 2400 / image.shape[1])
    small = cv2.resize(image, None, fx=scale, fy=scale, interpolation=cv2.INTER_AREA)
    width = small.shape[1]
    lines = cv2.HoughLinesP(binary(small), 1, np.pi / 1800,
                            threshold=max(50, int(width * .08)),
                            minLineLength=int(width * .20), maxLineGap=int(width * .025))
    if lines is None:
        raise ScanError("No long printed grid lines found. Include the entire score table.")
    return lines.reshape(-1, 4).astype(float) / scale


def deskew(image: np.ndarray) -> np.ndarray:
    lines = hough(image)
    dx, dy = lines[:, 2] - lines[:, 0], lines[:, 3] - lines[:, 1]
    angles = np.degrees(np.arctan2(dy, dx)) % 180
    weights = np.hypot(dx, dy)
    histogram, _ = np.histogram(angles, bins=180, range=(0, 180), weights=weights)
    # A circular histogram handles lines straddling 0 or 90 degrees.
    smoothed = sum(np.roll(histogram, k) for k in range(-4, 5))
    peak = float(np.argmax(smoothed))
    relative = (angles - peak + 90) % 180 - 90
    angle = (peak + np.median(relative[np.abs(relative) < 12]) + 90) % 180 - 90
    h, w = image.shape
    matrix = cv2.getRotationMatrix2D((w / 2, h / 2), float(angle), 1)
    new_w = int(np.ceil(abs(matrix[0, 0]) * w + abs(matrix[0, 1]) * h))
    new_h = int(np.ceil(abs(matrix[1, 0]) * w + abs(matrix[1, 1]) * h))
    matrix[:, 2] += ((new_w - w) / 2, (new_h - h) / 2)
    aligned = cv2.warpAffine(image, matrix, (new_w, new_h), borderValue=255)
    ys, xs = np.where(aligned < 230)
    if len(xs):
        pad = max(20, int(min(aligned.shape) * .015))
        aligned = aligned[max(0, ys.min() - pad):min(new_h, ys.max() + pad + 1),
                          max(0, xs.min() - pad):min(new_w, xs.max() + pad + 1)]
    return aligned


def find_grid(image: np.ndarray) -> list[tuple[float, float]]:
    """Find all 41 actual boundaries; no fixed page fractions or guessed rows."""
    width = image.shape[1]
    lines = hough(image)
    records = []
    for x0, y0, x1, y1 in lines:
        if abs(x1 - x0) < 1:
            continue
        slope = (y1 - y0) / (x1 - x0)
        if abs(slope) <= .18:
            records.append((y0 + slope * (width / 2 - x0), (x0, y0, x1, y1)))
    groups: list[list] = []
    tolerance = max(3., image.shape[0] * .0035)
    for center, line in sorted(records):
        if groups and abs(center - np.median([a[0] for a in groups[-1]])) < tolerance:
            groups[-1].append((center, line))
        else:
            groups.append([(center, line)])
    centers, coefficients = [], []
    for group in groups:
        points = np.array([p for _, l in group for p in [(l[0], l[1]), (l[2], l[3])]])
        slope, intercept = np.polyfit(points[:, 0], points[:, 1], 1)
        centers.append(slope * width / 2 + intercept)
        coefficients.append((slope, intercept))
    order = np.argsort(centers)
    centers = np.array(centers)[order]
    coefficients = [coefficients[i] for i in order]
    gaps = np.diff(centers)
    plausible = gaps[(gaps > image.shape[0] * .008) & (gaps < image.shape[0] * .05)]
    if len(plausible) < 35:
        raise ScanError("Fewer than 41 grid boundaries found; the table is cropped or blurred.")
    step = float(np.median(plausible))
    candidates = []
    for start in range(len(centers)):
        sequence = [start]
        while True:
            last = sequence[-1]
            choices = [k for k in range(last + 1, len(centers))
                       if .68 * step < centers[k] - centers[last] < 1.38 * step]
            if not choices:
                break
            sequence.append(min(choices, key=lambda k: abs(centers[k] - centers[last] - step)))
        if len(sequence) >= 41:
            selected = sequence[:41]
            variation = np.std(np.diff(centers[selected])) / step
            if variation < .13:
                candidates.append((len(sequence) - variation, selected))
    if not candidates:
        raise ScanError("Could not identify exactly 40 consecutive target rows.")
    _, selected = max(candidates, key=lambda item: item[0])
    return [coefficients[i] for i in selected]


def remap(image: np.ndarray, xs: np.ndarray, ys: np.ndarray) -> np.ndarray:
    return cv2.remap(image, xs.astype(np.float32), ys.astype(np.float32),
                     cv2.INTER_CUBIC, borderMode=cv2.BORDER_CONSTANT, borderValue=255)


def trace_rule(image: np.ndarray, slope: float, intercept: float, step: float):
    """Follow local rule curvature to tolerate folded pages as well as skew."""
    width = image.shape[1]
    x = np.arange(width, dtype=float)
    center = slope * x + intercept
    radius = max(3, int(step * .26))
    offsets = np.arange(-radius, radius + 1)
    strip = remap(image, np.tile(x, (len(offsets), 1)), center[None, :] + offsets[:, None])
    dark = (strip < 165).astype(float)
    block = max(25, int(width * .025))
    anchor_x, anchor_offsets = [], []
    for left in range(0, width, block):
        signal = gaussian_filter1d(dark[:, left:left + block].mean(axis=1), .7)
        anchor_x.append(min(left + block / 2, width - 1))
        anchor_offsets.append(int(np.argmax(signal)) - radius)
    curve = center + np.interp(x, anchor_x, anchor_offsets)
    strip = remap(image, np.tile(x, (5, 1)), curve[None, :] + np.arange(-2, 3)[:, None])
    signal = gaussian_filter1d(np.mean(strip < 165, axis=0), 4)
    good = (signal > .2).astype(np.uint8).reshape(1, -1)
    good[:, :int(width * .008)] = 0
    good[:, int(width * .985):] = 0
    good = cv2.morphologyEx(good, cv2.MORPH_CLOSE, np.ones((1, max(3, int(width * .012))), np.uint8))
    indices = np.flatnonzero(np.diff(np.r_[0, good[0], 0].astype(int)))
    runs = list(zip(indices[::2], indices[1::2]))
    left, right = max(runs, key=lambda p: p[1] - p[0]) if runs else (0, width - 1)
    return curve, float(left), float(right - 1)


def straighten(image: np.ndarray, rules: list[tuple[float, float]]):
    width = image.shape[1]
    step = np.median(np.diff([m * width / 2 + b for m, b in rules]))
    traced = [trace_rule(image, m, b, step) for m, b in rules]
    curves = np.array([r[0] for r in traced])
    # Rule endpoints determine horizontal perspective; reject short fragments.
    ends = np.array([[r[1], r[2]] for r in traced])
    valid = (ends[:, 1] - ends[:, 0]) > .65 * width
    if np.count_nonzero(valid) < 25:
        raise ScanError("Not enough complete table rules to correct perspective.")
    row_numbers = np.arange(41)
    for column in range(2):
        fit = np.polyfit(row_numbers[valid], ends[valid, column], 1)
        residual = ends[:, column] - np.polyval(fit, row_numbers)
        inliers = valid & (np.abs(residual) < width * .025)
        if np.count_nonzero(inliers) < 20:
            raise ScanError("Table edges are inconsistent; cannot map lane columns safely.")
        ends[:, column] = np.polyval(np.polyfit(row_numbers[inliers], ends[inliers, column], 1), row_numbers)
    out_width = int(np.median(ends[:, 1] - ends[:, 0]))
    u = np.linspace(0, 1, out_width)
    row_x = ends[:, 0, None] + (ends[:, 1] - ends[:, 0])[:, None] * u
    row_y = np.array([np.interp(row_x[i], np.arange(width), curves[i]) for i in range(41)])
    if np.any(np.diff(row_y, axis=0) < step * .45):
        raise ScanError("Grid boundaries overlap; rescan with the entire table in focus.")
    v = np.arange(40 * ROW_HEIGHT) / ROW_HEIGHT
    row = np.minimum(v.astype(int), 39)
    alpha = (v - row)[:, None]
    grid = remap(image, row_x[row] * (1 - alpha) + row_x[row + 1] * alpha,
                 row_y[row] * (1 - alpha) + row_y[row + 1] * alpha)
    # Header and printed totals are sampled outside, independently of the hits.
    top_step, bottom_step = row_y[1] - row_y[0], row_y[-1] - row_y[-2]
    header_offsets = np.linspace(-1.12, -.10, 56)
    header = remap(image, np.tile(row_x[0], (56, 1)),
                   row_y[0][None, :] + header_offsets[:, None] * top_step[None, :])
    total_offsets = np.linspace(.09, 1.18, 62)
    totals = remap(image, np.tile(row_x[-1], (62, 1)),
                   row_y[-1][None, :] + total_offsets[:, None] * bottom_step[None, :])
    return grid, header, totals


def remove_rules(image: np.ndarray) -> np.ndarray:
    ink = binary(image)
    rules = cv2.morphologyEx(ink, cv2.MORPH_OPEN, np.ones((1, 55), np.uint8))
    return cv2.subtract(ink, rules)


def score_columns(grid: np.ndarray) -> list[int]:
    ink = remove_rules(grid)
    profile = gaussian_filter1d(np.count_nonzero(ink, axis=0).astype(float), 3)
    peaks, _ = find_peaks(profile, prominence=70, distance=max(30, int(grid.shape[1] * .023)))
    # A score column contains a repeated small glyph in most of the 40 rows.
    selected = []
    for x in peaks:
        if x < grid.shape[1] * .12:
            continue
        occupied = sum(np.count_nonzero(ink[r * ROW_HEIGHT + 7:(r + 1) * ROW_HEIGHT - 7,
                                              max(0, x - 15):x + 16]) >= 10 for r in range(40))
        if occupied >= 28:
            selected.append(int(x))
    if not 2 <= len(selected) <= 30:
        raise ScanError(f"Found {len(selected)} score columns; expected 2–30 lanes.")
    return selected


def read_bit(crop: np.ndarray) -> int | None:
    """Recognize only a narrow 1 or an oval 0, at two contrast thresholds."""
    answers = []
    threshold, _ = cv2.threshold(crop, 0, 255, cv2.THRESH_BINARY_INV | cv2.THRESH_OTSU)
    for threshold in (min(235, threshold + 5), min(245, threshold + 22)):
        ink = np.uint8(crop < threshold) * 255
        n, labels, stats, _ = cv2.connectedComponentsWithStats(ink, 8)
        components = []
        for i, (x, y, w, h, area) in enumerate(stats[1:], 1):
            if 9 <= h <= 36 and 2 <= w <= 30 and area >= 10:
                components.append((i, x, y, w, h))
        if len(components) != 1:
            return None
        i, x, y, w, h = components[0]
        glyph = np.uint8(labels[y:y + h, x:x + w] == i) * 255
        glyph = cv2.copyMakeBorder(glyph, 1, 1, 1, 1, cv2.BORDER_CONSTANT, value=0)
        contours, hierarchy = cv2.findContours(glyph, cv2.RETR_CCOMP, cv2.CHAIN_APPROX_SIMPLE)
        holes = sum(1 for k, c in enumerate(contours)
                    if hierarchy[0, k, 3] >= 0 and cv2.contourArea(c) >= 2) if hierarchy is not None else 0
        ratio = w / h
        if ratio < .44 and holes == 0:
            answers.append(1)
        elif .42 <= ratio <= 1.0 and holes == 1:
            answers.append(0)
        else:
            return None
    return answers[0] if answers[0] == answers[1] else None


def extract_hits(grid: np.ndarray, xs: list[int]) -> list[list[int | None]]:
    # Extract each lane independently. A total never changes any cell value.
    results, cells = [], []
    for j, x in enumerate(xs):
        pitch = min([abs(x - xs[k]) for k in (j - 1, j + 1) if 0 <= k < len(xs)])
        half = max(20, int(pitch * .40))
        hits, lane_cells = [], []
        for r in range(40):
            cell = grid[r * ROW_HEIGHT + 5:(r + 1) * ROW_HEIGHT - 5,
                        max(0, x - half):min(grid.shape[1], x + half + 1)]
            hits.append(read_bit(cell))
            lane_cells.append(cell)
        results.append(hits)
        cells.append(lane_cells)
    models = learn_digit_shapes(cells, results)
    if models:
        for j, hits in enumerate(results):
            for r, value in enumerate(hits):
                if value is None:
                    hits[r] = match_digit(cells[j][r], models)
    return results


def learn_digit_shapes(cells, readings) -> dict[int, np.ndarray]:
    """Learn this page's font from clear glyphs; no totals or external examples."""
    examples = {0: [], 1: []}
    for column, values in zip(cells, readings):
        for crop, value in zip(column, values):
            if value is None:
                continue
            threshold, _ = cv2.threshold(crop, 0, 255, cv2.THRESH_BINARY_INV | cv2.THRESH_OTSU)
            _, _, stats, _ = cv2.connectedComponentsWithStats(np.uint8(crop < threshold + 10) * 255, 8)
            glyphs = [(x, y, w, h) for x, y, w, h, area in stats[1:]
                      if 9 <= h <= 36 and 2 <= w <= 30 and area >= 10]
            if len(glyphs) != 1:
                continue
            x, y, w, h = glyphs[0]
            new_width = max(1, round(24 * w / h))
            if new_width > 26:
                continue
            glyph = cv2.resize(crop[y:y + h, x:x + w], (new_width, 24))
            canvas = np.full((32, 28), 255, np.uint8)
            left = (28 - new_width) // 2
            canvas[4:28, left:left + new_width] = glyph
            ink = (255 - canvas).astype(np.float32)
            ink /= max(1., float(np.percentile(ink, 99)))
            examples[value].append(ink)
    if any(len(examples[value]) < 8 for value in (0, 1)):
        return {}
    return {value: np.median(glyphs, axis=0).astype(np.float32) for value, glyphs in examples.items()}


def match_digit(crop: np.ndarray, models: dict[int, np.ndarray]) -> int | None:
    """Require two glyph regions to prefer the same digit by a clear margin.

    The upper regions tolerate a handwritten circle touching a glyph's foot.
    Both the absolute fit and separation from the other digit must pass.
    """
    ink = (255 - crop).astype(np.float32) / 255
    integral = cv2.integral(ink)
    answers = []
    for fraction, minimum, margin in ((.78, .69, .10), (.63, .68, .18)):
        scores = []
        for value in (0, 1):
            best = -1.
            for scale in (.8, .9, 1., 1.1):
                model = cv2.resize(models[value], None, fx=scale, fy=scale)
                model = model[:round(model.shape[0] * fraction)]
                height, width = model.shape
                if height > ink.shape[0] or width > ink.shape[1]:
                    continue
                score = cv2.matchTemplate(ink, model, cv2.TM_CCOEFF_NORMED)
                sums = integral[height:, width:] - integral[:-height, width:] - integral[height:, :-width] + integral[:-height, :-width]
                # Constant white regions can spuriously return correlation 1.
                score[sums < 8] = -1
                centers = np.arange(score.shape[1]) + width / 2
                score[:, np.abs(centers - crop.shape[1] / 2) > min(22, crop.shape[1] * .25)] = -1
                best = max(best, float(np.max(score)))
            scores.append(best)
        preferred = int(np.argmax(scores))
        if scores[preferred] < minimum or scores[preferred] - scores[1 - preferred] < margin:
            return None
        answers.append(preferred)
    return answers[0] if answers[0] == answers[1] else None


def clean_printed(crop: np.ndarray, *, total: bool) -> np.ndarray:
    ink = binary(crop)
    n, labels, stats, _ = cv2.connectedComponentsWithStats(ink, 8)
    clean = np.zeros_like(ink)
    for i, (x, y, w, h, area) in enumerate(stats[1:], 1):
        # Handwritten circles and names are taller than the printed characters.
        if 9 <= h <= crop.shape[0] * .68 and 2 <= w <= 38 and area >= 8:
            if y > 1 and y + h < crop.shape[0] - 1:
                clean[labels == i] = 255
    return 255 - clean


def tesseract(image: np.ndarray, psm: int = 7, whitelist: str | None = None) -> str:
    padded = cv2.copyMakeBorder(image, 12, 12, 12, 12, cv2.BORDER_CONSTANT, value=255)
    enlarged = cv2.resize(padded, None, fx=3, fy=3, interpolation=cv2.INTER_CUBIC)
    encoded = cv2.imencode(".png", enlarged)[1].tobytes()
    command = ["tesseract", "stdin", "stdout", "--psm", str(psm)]
    if whitelist:
        command += ["-c", "tessedit_char_whitelist=" + whitelist]
    environment = dict(os.environ, OMP_THREAD_LIMIT="1")
    result = subprocess.run(command, input=encoded, stdout=subprocess.PIPE,
                            stderr=subprocess.PIPE, timeout=20, env=environment)
    if result.returncode:
        raise RuntimeError("Tesseract failed: " + result.stderr.decode(errors="replace")[:300])
    return result.stdout.decode(errors="replace").strip()


def read_number(crop: np.ndarray, *, lane: bool) -> int | None:
    # Keep grayscale strokes in the printed baseline band. A binary mask can
    # accidentally join a 7 to a handwritten circle or turn an angled 1 into 7.
    raw_readings = []
    for lower in (.08, .16):
        band = crop[int(crop.shape[0] * lower):int(crop.shape[0] * .78)]
        for psm in (7, 6):
            text = re.sub(r"\s+", "", tesseract(band, psm, "L0123456789" if lane else "0123456789"))
            match = re.fullmatch(r"L(\d{1,2})" if lane else r"(\d{1,2})", text)
            if match:
                value = int(match[1])
                if (1 <= value <= 30) if lane else (0 <= value <= 40):
                    raw_readings.append(value)
    raw_counts = Counter(raw_readings)
    if raw_counts:
        value, count = raw_counts.most_common(1)[0]
        if count >= 3 and len(raw_counts) == 1:
            return value
    cleaned = clean_printed(crop, total=not lane)
    readings = []
    if lane:
        yy, xx = np.where(cleaned < 128)
        if len(xx):
            full_label = cleaned[yy.min():yy.max() + 1, xx.min():xx.max() + 1]
            for psm in (7, 6):
                raw = re.sub(r"\s+", "", tesseract(full_label, psm, "L0123456789"))
                match = re.fullmatch(r"L(\d{1,2})", raw)
                if match and 1 <= int(match[1]) <= 30:
                    readings.append(int(match[1]))
        # Drop an actual L-shaped prefix, rather than treating an OCR '1' as L.
        # A prefix joined to handwriting may already have been removed above.
        ink = np.uint8(cleaned < 128) * 255
        _, components, stats, _ = cv2.connectedComponentsWithStats(ink, 8)
        glyphs = sorted(enumerate(stats[1:], 1), key=lambda item: item[1][0])
        if len(glyphs) > 1:
            i, (x, y, w, h, _) = glyphs[0]
            glyph = components[y:y + h, x:x + w] == i
            top_width = np.mean(np.sum(glyph[:max(1, h // 2)], axis=1))
            foot_width = np.max(np.sum(glyph[max(1, h * 3 // 4):], axis=1))
            if w / h > .3 and foot_width > 1.9 * top_width:
                cleaned[components == i] = 255
    if not np.any(cleaned < 128):
        return None
    yy, xx = np.where(cleaned < 128)
    cleaned = cleaned[yy.min():yy.max() + 1, xx.min():xx.max() + 1]
    # The isolated printed zero is unambiguous by its single interior hole.
    padded = cv2.copyMakeBorder(cleaned, 5, 5, 10, 10, cv2.BORDER_CONSTANT, value=255)
    bit = read_bit(padded)
    for psm in (7, 6):
        raw = re.sub(r"\s+", "", tesseract(cleaned, psm, "0123456789"))
        match = re.fullmatch(r"(\d{1,2})", raw)
        if match:
            value = int(match[1])
            if (1 <= value <= 30) if lane else (0 <= value <= 40):
                readings.append(value)
    counts = Counter(readings)
    if not lane and readings == [0] and bit == 0:
        return 0
    agreed = [value for value, count in counts.items() if count >= 2]
    if len(agreed) == 1 and all(value == agreed[0] for value in raw_readings):
        return agreed[0]
    return None


def check_lane_order(labels: list[int | None]) -> list[int | None]:
    """Reject out-of-order OCR; retain only unambiguous increasing anchors.

    Gaps are allowed. Never fill a missing label from its neighbours.
    """
    anchors = [(i, value) for i, value in enumerate(labels) if value is not None]
    if not anchors:
        return labels.copy()
    forward, backward = [], [1] * len(anchors)
    for j, (_, value) in enumerate(anchors):
        forward.append(1 + max((forward[k] for k in range(j) if anchors[k][1] < value), default=0))
    for j in reversed(range(len(anchors))):
        backward[j] = 1 + max((backward[k] for k in range(j + 1, len(anchors)) if anchors[k][1] > anchors[j][1]), default=0)
    longest = max(forward)
    levels = Counter(forward[j] for j in range(len(anchors)) if forward[j] + backward[j] - 1 == longest)
    result = [None] * len(labels)
    for j, (index, value) in enumerate(anchors):
        if forward[j] + backward[j] - 1 == longest and levels[forward[j]] == 1:
            result[index] = value
    return result


def validate_lane(lane: Lane) -> None:
    if lane.label is None:
        lane.problems.append("Printed lane label is unreadable; column number is not a lane number")
    missing = [str(i + 1) for i, value in enumerate(lane.hits) if value is None]
    if missing:
        lane.problems.append("Unreadable hit rows: " + ", ".join(missing))
    if lane.printed_total is None:
        lane.problems.append("Printed total is unreadable")
    if lane.difference is not None and lane.difference != 0:
        lane.problems.append(f"Total mismatch: extracted {lane.calculated_total}, printed {lane.printed_total}; check all 40 rows")


def read_distances(grid: np.ndarray, first_column: int) -> list[int | None]:
    """Read the actual printed row distances instead of trusting legacy headers."""
    def one_row(row):
        crop = grid[row * ROW_HEIGHT + 7:(row + 1) * ROW_HEIGHT - 7,
                    :max(100, first_column - 60)]
        readings = []
        for psm in (7, 6):
            text = tesseract(crop, psm)
            match = re.search(r"(50|100|150|200|250|300)\s*[mM]", text)
            if match:
                readings.append(int(match[1]))
        return readings[0] if len(readings) == 2 and readings[0] == readings[1] else None
    with ThreadPoolExecutor(max_workers=4) as pool:
        return list(pool.map(one_row, range(40)))


def extract_page(image: np.ndarray, source: str, page: int,
                 debug_dir: Path | None = None, workers: int = 4) -> Sheet:
    aligned = deskew(image)
    for attempt in range(2):
        try:
            rules = find_grid(aligned)
            grid, header, totals = straighten(aligned, rules)
            xs = score_columns(grid)
            crops = []
            for j, x in enumerate(xs):
                pitch = min([abs(x - xs[k]) for k in (j - 1, j + 1) if 0 <= k < len(xs)])
                half = int(pitch * .43)
                crops.append((header[:, max(0, x - half):x + half], totals[:, max(0, x - half):x + half]))
            with ThreadPoolExecutor(max_workers=workers) as pool:
                labels = list(pool.map(lambda c: read_number(c[0], lane=True), crops))
            if sum(label is not None for label in labels) < max(2, len(xs) // 3):
                raise ScanError("Cannot confirm printed lane labels or page orientation.")
            break
        except ScanError:
            if attempt:
                raise
            aligned = cv2.rotate(aligned, cv2.ROTATE_180)
    with ThreadPoolExecutor(max_workers=workers) as pool:
        printed = list(pool.map(lambda c: read_number(c[1], lane=False), crops))
    labels = check_lane_order(labels)
    hits = extract_hits(grid, xs)
    lanes = [Lane(f"L{label}" if label is not None else None, bits, total, c[0], c[1])
             for label, bits, total, c in zip(labels, hits, printed, crops)]
    duplicates = [label for label, count in Counter(labels).items() if label is not None and count > 1]
    for lane in lanes:
        if lane.label and int(lane.label[1:]) in duplicates:
            lane.problems.append("Duplicate lane label; numbering requires review")
        validate_lane(lane)
    distances = read_distances(grid, xs[0])
    problems = []
    if 1 not in labels:
        problems.append("Lane 1 could not be identified; check the complete left edge of the scan")
    if any(d is None for d in distances):
        problems.append("Printed target distances could not all be read; target numbers are preserved")
    # Printed session metadata is optional and never supplies numeric scores.
    first_y = min(m * aligned.shape[1] / 2 + b for m, b in rules)
    metadata = tesseract(aligned[:max(1, int(first_y * .45)), :], 6)
    date = re.search(r"\b(\d{1,2})\s+([A-Za-z]{3})\s+(20\d{2})\s*[-–]\s*(\d{1,2}:\d{2})", metadata)
    firing = re.search(r"Firing\s+Order\s*[:;]?\s*(\d+)", metadata, re.I)
    session = f"{int(date[1]):02d} {date[2].title()} {date[3]} • {date[4]}" if date else ""
    sheet = Sheet(source, page, lanes, distances, session, int(firing[1]) if firing else None, problems)
    if debug_dir:
        debug_dir.mkdir(parents=True, exist_ok=True)
        # Filename hashes keep same-named PDFs from different directories separate.
        import hashlib
        stem = re.sub(r"[^\w.-]", "_", Path(source).stem) + "-" + hashlib.sha256(source.encode()).hexdigest()[:8] + f"-p{page}"
        visual = cv2.cvtColor(grid, cv2.COLOR_GRAY2BGR)
        for j, x in enumerate(xs):
            cv2.line(visual, (x, 0), (x, grid.shape[0] - 1), (0, 100, 255), 1)
            for r, value in enumerate(hits[j]):
                if value is None:
                    cv2.rectangle(visual, (x - 25, r * ROW_HEIGHT), (x + 25, (r + 1) * ROW_HEIGHT), (0, 0, 255), 2)
        cv2.imwrite(str(debug_dir / (stem + "-grid.png")), visual)
        cv2.imwrite(str(debug_dir / (stem + "-labels.png")), header)
        cv2.imwrite(str(debug_dir / (stem + "-totals.png")), totals)
        report = {"source": source, "page": page, "problems": problems, "lanes": [
            {"column": i + 1, "lane": lane.label, "hits": lane.hits, "calculated_total": lane.calculated_total,
             "printed_total": lane.printed_total, "status": lane.status, "problems": lane.problems}
            for i, lane in enumerate(lanes)]}
        (debug_dir / (stem + ".json")).write_text(json.dumps(report, indent=2), encoding="utf-8")
    return sheet


def write_workbook(sheets: list[Sheet], output: Path, errors: list[tuple[str, int, str]],
                   template: Path | None = None) -> None:
    workbook = load_workbook(template) if template else Workbook()
    data = workbook.worksheets[0]
    common = sheets[0].distances if sheets else [None] * 40
    if any(sheet.distances != common for sheet in sheets):
        common = [None] * 40
    headers = ["Qualification Form", "Session / Source", "Firing Order", "Rank", "DoD ID Number", "Name", "Target / range"]
    headers += [f"{i:02d} · {distance} m" if distance is not None else f"{i:02d} · target"
                for i, distance in enumerate(common, 1)]
    headers += ["Calculated hits", "Printed total", "Check (0 = match)", "Status", "Source page", "Unreadable cells"]
    if template:
        if [data.cell(1, c).value for c in range(1, 8)] != headers[:7]:
            raise ValueError("Template must have the rangeCardData one-lane-per-row layout in its first sheet.")
        for c in range(8, 48):
            if not re.match(rf"^0?{c-7}\s*[·.:–-]", str(data.cell(1, c).value or "")):
                raise ValueError("Template target columns must be ordered 01–40 in H–AU.")
            match = re.search(r"(50|100|150|200|250|300)\s*m", str(data.cell(1, c).value or ""))
            if sheets and (match is None or common[c - 8] is None or int(match[1]) != common[c - 8]):
                raise ValueError("Template target distances disagree with the scan. Use a fresh workbook (omit --template) to preserve printed row order.")
        # Existing headers and records are preserved; append metadata in new columns.
        for c in range(51, 54):
            if data.cell(1, c).value is not None:
                raise ValueError(f"Template column {get_column_letter(c)} already contains data.")
        for c in range(51, 54):
            data.cell(1, c, headers[c - 1])
    else:
        data.title = "Sheet1"
        data.append(headers)
    review = workbook.create_sheet("Scan Review")
    review.append(["Source PDF", "Page", "Column", "Lane", "Calculated hits", "Printed total", "Difference", "Status", "Issue", "Lane label image", "Printed total image"])
    red = PatternFill("solid", fgColor="FCE4D6")
    green = PatternFill("solid", fgColor="E2F0D9")
    target_rows = workbook.create_sheet("Printed Target Rows")
    target_rows.append(["Source PDF", "Page", *[f"Target {i:02d} distance (m)" for i in range(1, 41)]])
    for sheet in sheets:
        target_rows.append([sheet.source, sheet.page, *sheet.distances])
        stamp = " • ".join(s for s in [sheet.session, f"Firing Order {sheet.firing_order}" if sheet.firing_order is not None else "",
                                      f"Source: {Path(sheet.source).name}, page {sheet.page}"] if s)
        for column, lane in enumerate(sheet.lanes, 1):
            row_status = "REVIEW" if lane.problems or sheet.problems else "MATCH"
            data.append([TITLE, stamp, sheet.firing_order, None, None, None, lane.label or f"UNREADABLE COLUMN {column}",
                         *lane.hits, lane.calculated_total, lane.printed_total, lane.difference, row_status, sheet.page, lane.hits.count(None)])
            row = data.max_row
            for c in (48, 49, 50, 51):
                data.cell(row, c).fill = green if row_status == "MATCH" else red
            for target, value in enumerate(lane.hits, 1):
                if value is None:
                    data.cell(row, target + 7).fill = red
                    data.cell(row, target + 7).comment = Comment("Unreadable scan cell; blank is not a miss (0).", "rangeFill")
            actual_distances = ", ".join(f"{i}: {d if d is not None else '?'} m" for i, d in enumerate(sheet.distances, 1))
            data.cell(row, 7).comment = Comment("Handwritten names are left for manual entry.\nPrinted target distances in source row order: " + actual_distances, "rangeFill")
            review.append([sheet.source, sheet.page, column, lane.label, lane.calculated_total,
                           lane.printed_total, lane.difference, row_status, "; ".join(lane.problems + sheet.problems)])
            review_row = review.max_row
            review.row_dimensions[review_row].height = 65
            for crop, anchor in ((lane.header_crop, "J"), (lane.total_crop, "K")):
                buffer = BytesIO(cv2.imencode(".png", crop)[1].tobytes())
                picture = ExcelImage(buffer)
                picture.height = 60
                picture.width = min(150, 60 * crop.shape[1] / crop.shape[0])
                review.add_image(picture, f"{anchor}{review_row}")
    for source, page, message in errors:
        review.append([source, page, None, None, None, None, None, "ERROR", message])
    for sheet in (data, review):
        sheet.freeze_panes = "H2" if sheet is data else "D2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="203864")
            cell.alignment = Alignment(wrap_text=True)
        sheet.row_dimensions[1].height = 32
    for c in range(8, 48):
        data.column_dimensions[get_column_letter(c)].width = 10
    for col, width in {"A": 24, "B": 65, "C": 15, "D": 12, "E": 20, "F": 25, "G": 25, "AV": 18, "AW": 18, "AX": 22, "AY": 16, "AZ": 15, "BA": 20}.items():
        data.column_dimensions[col].width = width
    for col, width in {"A": 45, "B": 8, "C": 8, "D": 12, "E": 18, "F": 18, "G": 12, "H": 12, "I": 85, "J": 23, "K": 23}.items():
        review.column_dimensions[col].width = width
    for row in review.iter_rows(min_row=2):
        row[8].alignment = Alignment(wrap_text=True, vertical="top")
    # No accidental overwrite; save atomically without damaging an existing file.
    output.parent.mkdir(parents=True, exist_ok=True)
    import tempfile
    with tempfile.NamedTemporaryFile(suffix=".xlsx", dir=output.parent, delete=False) as temporary:
        temp_path = Path(temporary.name)
    try:
        workbook.save(temp_path)
        os.link(temp_path, output)
    finally:
        temp_path.unlink(missing_ok=True)


def collect_inputs(paths: list[Path]) -> list[Path]:
    files = []
    for path in paths:
        if path.is_dir():
            files.extend(sorted(p for p in path.iterdir() if p.is_file() and p.suffix.lower() == ".pdf"))
        elif path.is_file() and path.suffix.lower() == ".pdf":
            files.append(path)
        else:
            raise ValueError(f"Not a PDF file or directory: {path}")
    return list(dict.fromkeys(p.resolve() for p in files))


def main(argv: list[str] | None = None) -> int:
    parser = ConverterParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("inputs", type=Path, nargs="+", help="PDF files or directories of PDFs")
    parser.add_argument("-o", "--output", type=Path, default=Path("range_scores.xlsx"))
    parser.add_argument("--template", type=Path, help="Append to a COPY of an existing rangeCardData workbook")
    parser.add_argument("--debug-dir", type=Path, help="Save corrected grid images and JSON extraction reports")
    parser.add_argument("--width", type=int, default=3000, help="Rasterized page width (default: 3000)")
    parser.add_argument("--workers", type=int, default=4, help="Concurrent local OCR calls (1–8)")
    args = parser.parse_args(argv)
    if not shutil.which("tesseract"):
        parser.error("Tesseract is missing. Install tesseract-ocr (Linux) or brew install tesseract (macOS).")
    if args.output.exists():
        parser.error(f"Output already exists; choose a new filename: {args.output}")
    if not 1800 <= args.width <= 6000 or not 1 <= args.workers <= 8:
        parser.error("--width must be 1800–6000 and --workers must be 1–8")
    try:
        files = collect_inputs(args.inputs)
        if not files:
            parser.error("No PDFs found")
    except ValueError as exc:
        parser.error(str(exc))
    sheets, errors = [], []
    for path in files:
        try:
            with fitz.open(path) as document:
                if document.needs_pass:
                    raise ValueError("PDF is password protected")
                if not document.page_count:
                    raise ValueError("PDF has no pages")
                for index, page in enumerate(document):
                    try:
                        zoom = args.width / page.rect.width
                        if args.width * page.rect.height * zoom > 45_000_000:
                            raise ScanError("Page dimensions exceed the safe rasterization limit")
                        pixmap = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), colorspace=fitz.csGRAY, alpha=False)
                        image = np.frombuffer(pixmap.samples, np.uint8).reshape(pixmap.height, pixmap.width).copy()
                        sheet = extract_page(image, str(path), index + 1, args.debug_dir, args.workers)
                        sheets.append(sheet)
                        matches = sum(lane.status == "MATCH" for lane in sheet.lanes)
                        print(f"{path.name}, page {index + 1}: {len(sheet.lanes)} lanes, {matches} numeric matches", flush=True)
                    except (ScanError, RuntimeError, subprocess.TimeoutExpired, cv2.error) as exc:
                        errors.append((str(path), index + 1, str(exc)))
                        print(f"ERROR {path.name}, page {index + 1}: {exc}", file=sys.stderr, flush=True)
        except (ValueError, RuntimeError, OSError) as exc:
            errors.append((str(path), 0, str(exc)))
            print(f"ERROR {path.name}: {exc}", file=sys.stderr)
    try:
        write_workbook(sheets, args.output, errors, args.template)
    except (ValueError, OSError) as exc:
        print(f"Cannot write workbook: {exc}", file=sys.stderr)
        return 1
    print(f"Saved {args.output}; inspect the Scan Review sheet.")
    if not sheets:
        return 1
    if errors or any(sheet.problems or any(lane.problems for lane in sheet.lanes) for sheet in sheets):
        print("REVIEW REQUIRED: unclear values remain blank; no scores were adjusted to match totals.")
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
